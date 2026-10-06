from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def create_excel_response(
    headers: list[str],
    rows: list[list],
    sheet_name: str = "Report",
):
    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = sheet_name

    # -----------------------------------------
    # Header
    # -----------------------------------------

    for column_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=header,
        )

        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )
        cell.alignment = Alignment(horizontal="center")

    # -----------------------------------------
    # Data
    # -----------------------------------------

    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(
            row,
            start=1,
        ):
            worksheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )

    # -----------------------------------------
    # Freeze header
    # -----------------------------------------

    worksheet.freeze_panes = "A2"

    # -----------------------------------------
    # Auto filter
    # -----------------------------------------

    if headers:
        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers))}" f"{len(rows) + 1}"
        )

    # -----------------------------------------
    # Column width
    # -----------------------------------------

    for column_index in range(
        1,
        len(headers) + 1,
    ):
        column_letter = get_column_letter(column_index)

        max_length = len(
            str(
                worksheet.cell(
                    row=1,
                    column=column_index,
                ).value
            )
        )

        for cell in worksheet[column_letter]:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value)),
                )

        worksheet.column_dimensions[column_letter].width = min(
            max_length + 2,
            40,
        )

    # -----------------------------------------
    # Save to memory
    # -----------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output.getvalue()
